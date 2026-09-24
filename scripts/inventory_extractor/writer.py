"""Écriture du résultat en CSV ou Excel."""
from __future__ import annotations

import csv

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension


def write_csv(rows: list[dict], columns: list[str], path: str) -> None:
    # utf-8-sig : Excel ouvre correctement les accents sans manipulation
    # d'encodage manuelle de la part de l'utilisateur.
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict], columns: list[str], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hosts"

    ws.append(columns)
    for row in rows:
        ws.append([row.get(col, "") for col in columns])

    # Entêtes figées (ligne 1) + auto-filtre sur toute la plage utilisée,
    # pour que l'utilisateur trie/filtre lui-même dans Excel (acté en 4.3).
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Largeur de colonne auto, plafonnée pour éviter des colonnes démesurées
    # sur des champs texte libres (ex: notes, software_full).
    for idx, column in enumerate(columns, start=1):
        max_len = max(
            [len(column)] + [len(str(row.get(column, ""))) for row in rows]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 60)

    wb.save(path)
